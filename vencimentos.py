"""Controle de vencimentos de SST: lê controle-sst.xlsx e gera o relatório HTML.

Uso:
    py -3.13 vencimentos.py                          # planilha padrão, abre o relatório no navegador
    py -3.13 vencimentos.py --planilha D:/SST/x.xlsx # outra planilha
    py -3.13 vencimentos.py --hoje 2026-12-01        # simula outra data
    py -3.13 vencimentos.py --nao-abrir              # só gera o arquivo

Regras:
- Para cada funcionário ativo ou afastado, a aba Funções diz o que a função exige.
- Vale o registro mais recente de cada (matrícula, item). Vence = data + validade, ou a coluna "Vence em".
- Item exigido sem registro = FALTANDO. Item sem validade só precisa existir.
- Registro de item que a função não exige só aparece se estiver vencido ou perto de vencer.

O HTML é montado em relatorio.py.
"""
import argparse
import calendar
import datetime as dt
import sys
import unicodedata
import webbrowser
from collections import Counter
from pathlib import Path

import esquema as E

PASTA = Path(__file__).resolve().parent


def somar_meses(d, meses):
    ano, mes = divmod(d.month - 1 + meses, 12)
    ano += d.year
    mes += 1
    return dt.date(ano, mes, min(d.day, calendar.monthrange(ano, mes)[1]))


def normalizar(texto):
    texto = unicodedata.normalize("NFKD", str(texto or ""))
    return "".join(c for c in texto if not unicodedata.combining(c)).casefold().strip()


def para_data(valor):
    if valor is None or valor == "":
        return None
    if isinstance(valor, dt.datetime):
        return valor.date()
    if isinstance(valor, dt.date):
        return valor
    texto = str(valor).strip()
    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%y", "%d-%m-%Y"):
        try:
            return dt.datetime.strptime(texto, formato).date()
        except ValueError:
            pass
    raise ValueError(f"data inválida: {texto!r}")


def texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def ler_aba(wb, nome, colunas, obrigatoria=True):
    """Lê uma aba como lista de (nº da linha, {coluna: valor}), achando colunas pelo nome sem acento."""
    abas = {normalizar(ws.title): ws for ws in wb.worksheets}
    ws = abas.get(normalizar(nome))
    if ws is None:
        if obrigatoria:
            sys.exit(f"A planilha não tem a aba '{nome}'. Gere um modelo novo com criar_planilha.py.")
        return [], []
    linhas = ws.iter_rows(values_only=True)
    cab_original = list(next(linhas, []))
    cab = [normalizar(c) for c in cab_original]
    conhecidas = {normalizar(c) for c in colunas}
    faltam = [c for c in colunas if normalizar(c) not in cab]
    if faltam:
        sys.exit(f"A aba '{nome}' não tem as colunas: {', '.join(faltam)}.")
    pos = {c: cab.index(normalizar(c)) for c in colunas}
    extras = [(i, texto(cab_original[i])) for i, c in enumerate(cab) if c and c not in conhecidas]
    dados = []
    for n, linha in enumerate(linhas, start=2):
        if not linha or all(v in (None, "") for v in linha):
            continue
        registro = {c: (linha[p] if p < len(linha) else None) for c, p in pos.items()}
        registro["_extras"] = {nome_col: linha[i] if i < len(linha) else None for i, nome_col in extras}
        dados.append((n, registro))
    return dados, [nome_col for _, nome_col in extras]


def classificar(vence, hoje):
    if vence is None:
        return "ok", None
    dias = (vence - hoje).days
    if dias < 0:
        return "vencido", dias
    if dias <= E.ALERTA_CURTO:
        return "30", dias
    if dias <= E.ALERTA_LONGO:
        return "60", dias
    return "ok", dias


def analisar(wb, hoje):
    problemas = []

    def problema(aba, linha, msg):
        problemas.append(f"{aba}, linha {linha}: {msg}")

    # Itens
    itens = {}
    linhas, _ = ler_aba(wb, E.ABA_ITENS, E.COLS_ITENS)
    for n, r in linhas:
        cod = texto(r["Código"]).upper()
        if not cod:
            continue
        validade = r["Validade (meses)"]
        try:
            meses = int(float(validade)) if texto(validade) else None
        except ValueError:
            problema(E.ABA_ITENS, n, f"validade {validade!r} não é número; tratado como 'não vence'")
            meses = None
        itens[cod] = {"descricao": texto(r["Descrição"]) or cod, "meses": meses}

    # Funções: a 1ª coluna é a função, as demais são códigos de itens
    linhas, codigos_colunas = ler_aba(wb, E.ABA_FUNCOES, ["Função"])
    for cod in codigos_colunas:
        if cod.upper() not in itens:
            problemas.append(f"{E.ABA_FUNCOES}: a coluna '{cod}' não existe na aba {E.ABA_ITENS}")
    funcoes = {}
    for n, r in linhas:
        funcao = texto(r["Função"])
        if not funcao:
            continue
        exige = {}
        for cod, marca in r["_extras"].items():
            cod = cod.upper()
            if cod not in itens or texto(marca) == "":
                continue
            if texto(marca).upper() == "X":
                exige[cod] = itens[cod]["meses"]
            else:
                try:
                    exige[cod] = int(float(texto(marca).replace(",", ".")))
                except ValueError:
                    problema(E.ABA_FUNCOES, n, f"'{marca}' em {cod} não é X nem número; ignorado")
        funcoes[normalizar(funcao)] = exige

    # Funcionários
    funcionarios = {}
    linhas, _ = ler_aba(wb, E.ABA_FUNCIONARIOS, E.COLS_FUNCIONARIOS)
    for n, r in linhas:
        mat = texto(r["Matrícula"])
        if not mat:
            problema(E.ABA_FUNCIONARIOS, n, "sem matrícula; linha ignorada")
            continue
        if mat in funcionarios:
            problema(E.ABA_FUNCIONARIOS, n, f"matrícula {mat} repetida; vale a primeira")
            continue
        situacao = texto(r["Situação"]) or "Ativo"
        if normalizar(situacao) not in map(normalizar, E.SITUACOES):
            problema(E.ABA_FUNCIONARIOS, n, f"situação '{situacao}' desconhecida; tratado como Ativo")
            situacao = "Ativo"
        funcao = texto(r["Função"])
        if normalizar(situacao) != "desligado" and normalizar(funcao) not in funcoes:
            problema(E.ABA_FUNCIONARIOS, n, f"função '{funcao}' de {mat} não está na aba {E.ABA_FUNCOES}; nada é exigido dela")
        funcionarios[mat] = {
            "nome": texto(r["Nome"]), "funcao": funcao, "obra": texto(r["Obra"]) or "(sem obra)",
            "empresa": texto(r["Empresa"]) or "Própria", "situacao": situacao,
        }

    # Registros: fica o mais recente de cada (matrícula, código)
    ultimo = {}
    linhas, _ = ler_aba(wb, E.ABA_REGISTROS, E.COLS_REGISTROS)
    for n, r in linhas:
        mat, cod = texto(r["Matrícula"]), texto(r["Código"]).upper()
        if mat not in funcionarios:
            problema(E.ABA_REGISTROS, n, f"matrícula {mat or '(vazia)'} não está em {E.ABA_FUNCIONARIOS}")
            continue
        if cod not in itens:
            problema(E.ABA_REGISTROS, n, f"código {cod or '(vazio)'} não está na aba {E.ABA_ITENS}")
            continue
        try:
            data = para_data(r["Data de realização"])
            vence = para_data(r["Vence em"])
        except ValueError as erro:
            problema(E.ABA_REGISTROS, n, str(erro))
            continue
        if data is None:
            problema(E.ABA_REGISTROS, n, "sem data de realização")
            continue
        if data > hoje:
            problema(E.ABA_REGISTROS, n, f"data de realização no futuro ({data:%d/%m/%Y})")
        chave = (mat, cod)
        if chave not in ultimo or data >= ultimo[chave]["data"]:
            ultimo[chave] = {"data": data, "vence": vence}

    # Situação de cada exigência
    linhas_rel = []
    for mat, f in funcionarios.items():
        if normalizar(f["situacao"]) == "desligado":
            continue
        exige = funcoes.get(normalizar(f["funcao"]), {})
        codigos = list(exige) + [cod for (m, cod) in ultimo if m == mat and cod not in exige]
        for cod in codigos:
            exigido = cod in exige
            meses = exige.get(cod, itens[cod]["meses"])
            reg = ultimo.get((mat, cod))
            notas = []
            if reg is None:
                status, dias, vence, data = "faltando", None, None, None
            else:
                data = reg["data"]
                if reg["vence"]:
                    vence = reg["vence"]
                    notas.append("vencimento informado na planilha")
                else:
                    vence = somar_meses(data, meses) if meses else None
                status, dias = classificar(vence, hoje)
            if not exigido:
                if status not in ("vencido", "30", "60"):
                    continue
                notas.append("a função não exige")
            if normalizar(f["situacao"]) == "afastado":
                notas.append("afastado")
            linhas_rel.append({
                **f, "matricula": mat, "codigo": cod, "item": itens[cod]["descricao"],
                "data": data, "vence": vence, "dias": dias, "status": status, "nota": "; ".join(notas),
            })
    linhas_rel.sort(key=lambda x: (E.STATUS[x["status"]][2], x["dias"] if x["dias"] is not None else 0, x["obra"], x["nome"]))

    # CA dos EPIs em uso
    cas = []
    linhas, _ = ler_aba(wb, E.ABA_CA, E.COLS_CA, obrigatoria=False)
    for n, r in linhas:
        if normalizar(r["Em uso"]) in ("nao", "n"):
            continue
        try:
            validade = para_data(r["Validade do CA"])
        except ValueError as erro:
            problema(E.ABA_CA, n, str(erro))
            continue
        if validade is None:
            problema(E.ABA_CA, n, "CA sem data de validade")
            continue
        status, dias = classificar(validade, hoje)
        cas.append({"epi": texto(r["EPI"]), "fabricante": texto(r["Fabricante"]), "ca": texto(r["CA"]),
                    "vence": validade, "dias": dias, "status": status})
    cas.sort(key=lambda x: x["dias"])

    return {"linhas": linhas_rel, "cas": cas, "problemas": problemas,
            "pessoas": sum(1 for f in funcionarios.values() if normalizar(f["situacao"]) != "desligado")}


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description="Gera o relatório de vencimentos de SST.")
    ap.add_argument("--planilha", default=str(PASTA / "controle-sst.xlsx"))
    ap.add_argument("--hoje", help="data de referência AAAA-MM-DD (padrão: hoje)")
    ap.add_argument("--saida", help="arquivo HTML (padrão: relatorios/vencimentos-AAAA-MM-DD.html)")
    ap.add_argument("--nao-abrir", action="store_true", help="não abre o navegador")
    args = ap.parse_args()

    from openpyxl import load_workbook
    from relatorio import gerar_html

    planilha = Path(args.planilha)
    if not planilha.exists():
        sys.exit(f"Não achei {planilha}. Crie o modelo com: py -3.13 criar_planilha.py")
    hoje = dt.date.fromisoformat(args.hoje) if args.hoje else dt.date.today()

    try:
        wb = load_workbook(planilha, data_only=True, read_only=True)
    except PermissionError:
        sys.exit(f"Não consegui abrir {planilha}. Feche a planilha no Excel e rode de novo.")
    res = analisar(wb, hoje)
    wb.close()

    saida = Path(args.saida) if args.saida else PASTA / "relatorios" / f"vencimentos-{hoje:%Y-%m-%d}.html"
    saida.parent.mkdir(parents=True, exist_ok=True)
    saida.write_text(gerar_html(res, hoje, planilha), encoding="utf-8")

    cont = Counter(l["status"] for l in res["linhas"])
    ca = Counter(c["status"] for c in res["cas"])
    print(f"Vencimentos de SST em {hoje:%d/%m/%Y} ({res['pessoas']} pessoas)")
    for s in ("vencido", "faltando", "30", "60", "ok"):
        print(f"  {E.STATUS[s][1]:<28} {cont.get(s, 0):>4}")
    print(f"  {'CA vencido / a vencer':<28} {ca.get('vencido', 0):>4} / {ca.get('30', 0) + ca.get('60', 0)}")
    if res["problemas"]:
        print(f"  {len(res['problemas'])} problema(s) na planilha: veja o fim do relatório")
    print(f"Relatório: {saida}")
    if not args.nao_abrir:
        webbrowser.open(saida.as_uri())


if __name__ == "__main__":
    main()
