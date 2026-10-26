"""Cria o modelo controle-sst.xlsx do controle de vencimentos de SST.

Uso:
    py -3.13 criar_planilha.py                 # modelo vazio, só com itens e funções padrão
    py -3.13 criar_planilha.py --exemplo       # com funcionários e registros fictícios para testar
    py -3.13 criar_planilha.py --saida X.xlsx  # outro nome/lugar

Não sobrescreve uma planilha existente, a não ser com --forcar.
"""
import argparse
import datetime as dt
import random
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import esquema as E
from vencimentos import somar_meses

PASTA = Path(__file__).resolve().parent
COR_CABECALHO = PatternFill("solid", fgColor="1F3A5F")
FONTE_CABECALHO = Font(bold=True, color="FFFFFF")
FORMATO_DATA = "DD/MM/YYYY"

LEIAME = [
    "CONTROLE DE VENCIMENTOS DE SST",
    "",
    "Como preencher",
    "1. Itens: o que é controlado (ASO, treinamentos NR, OS...). Validade em meses; vazio = não vence, só precisa existir.",
    "2. Funções: uma linha por função, uma coluna por item. X = exigido com a validade padrão; número = exigido com essa validade em meses.",
    "3. Funcionários: um por linha, próprios e terceirizados. Empresa = 'Própria' ou o nome da terceirizada. Situação: Ativo, Afastado ou Desligado.",
    "4. Registros: cada ASO, treinamento ou OS feito vira uma linha. O script usa o registro mais recente de cada pessoa e item.",
    "   'Vence em' é opcional: preencha só quando o vencimento for diferente do calculado (ex.: o médico reduziu o prazo).",
    "5. EPI e CA: os CAs dos EPIs em uso, com a validade consultada no CAEPI.",
    "",
    "Como gerar o relatório",
    "Dê dois cliques em RODAR.bat (ou rode: py -3.13 vencimentos.py). O relatório abre no navegador e fica salvo em relatorios/.",
    "",
    "Cuidados",
    "- Os prazos padrão da aba Itens foram conferidos em 26/10/2026, mas quem manda é o PGR e o PCMSO da empresa: ajuste.",
    "- Não registre resultado de exame nem CID aqui. Só datas e aptidão (LGPD: dado de saúde é sensível).",
    "- Desligados não aparecem no relatório.",
]

NOMES = ["Ana", "Bruno", "Carla", "Diego", "Elaine", "Fábio", "Gabriel", "Helena", "Igor", "Joana",
         "Kleber", "Larissa", "Marcos", "Natália", "Otávio", "Paula", "Rafael", "Sandra", "Tiago", "Vanessa",
         "Wagner", "Yasmin", "Alex", "Beatriz", "Cícero", "Daniela", "Everton", "Fernanda", "Gilson", "Ivone"]
SOBRENOMES = ["Teste", "Exemplo", "Modelo", "Fictício", "Simulado", "Demo"]
OBRAS = ["Obra 01 - Residencial Ipê", "Obra 02 - Galpão Logístico", "Obra 03 - Reforma Escola"]
EMPRESAS = ["Própria"] * 4 + ["Empreiteira Alfa (terc.)", "Elétrica Beta (terc.)"]


def cabecalho(ws, colunas, larguras=None):
    ws.append(colunas)
    for i, _ in enumerate(colunas, start=1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font = COR_CABECALHO, FONTE_CABECALHO
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = (larguras or {}).get(i, 16)
    ws.freeze_panes = "A2"


def formatar_datas(ws, colunas, ate_linha=2000):
    for col in colunas:
        for linha in range(2, ate_linha + 1):
            ws.cell(row=linha, column=col).number_format = FORMATO_DATA


def lista(ws, intervalo, formula):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.error, dv.errorTitle = "Escolha um valor da lista.", "Valor inválido"
    ws.add_data_validation(dv)
    dv.add(intervalo)


def dados_exemplo(hoje, rng):
    """Funcionários e registros fictícios, espalhados entre vencido, 30 dias, 60 dias, faltando e em dia."""
    funcoes = list(E.FUNCOES_PADRAO)
    itens = {cod: meses for cod, _, meses, _, _ in E.ITENS_PADRAO}
    funcionarios, registros = [], []
    for i, nome in enumerate(NOMES):
        mat = f"{1001 + i}"
        funcao = funcoes[i % len(funcoes)]
        empresa = rng.choice(EMPRESAS) if funcao != "Técnico de segurança" else "Própria"
        if funcao == "Eletricista":
            empresa = "Elétrica Beta (terc.)"
        situacao = "Desligado" if i == 7 else "Afastado" if i == 12 else "Ativo"
        admissao = hoje - dt.timedelta(days=rng.randint(200, 2500))
        funcionarios.append([mat, f"{nome} {rng.choice(SOBRENOMES)}", funcao, OBRAS[i % 3], empresa, admissao, situacao])

        for cod, marca in E.FUNCOES_PADRAO[funcao].items():
            meses = marca if isinstance(marca, int) else itens[cod]
            sorteio = rng.random()
            if meses is None:
                if sorteio < 0.08:
                    continue  # faltando
                registros.append([mat, cod, hoje - dt.timedelta(days=rng.randint(30, 700)), None, ""])
                continue
            if sorteio < 0.05:
                continue  # faltando
            elif sorteio < 0.13:
                vence = hoje - dt.timedelta(days=rng.randint(1, 120))   # vencido
            elif sorteio < 0.23:
                vence = hoje + dt.timedelta(days=rng.randint(0, 30))    # até 30 dias
            elif sorteio < 0.30:
                vence = hoje + dt.timedelta(days=rng.randint(31, 60))   # até 60 dias
            else:
                vence = hoje + dt.timedelta(days=rng.randint(61, 30 * meses - 10))
            registros.append([mat, cod, somar_meses(vence, -meses), None, ""])

    # Erros de propósito, para mostrar a seção "Problemas na planilha"
    registros.append(["9999", "ASO", hoje - dt.timedelta(days=100), None, "matrícula que não existe (exemplo)"])
    registros.append(["1001", "NR99", hoje - dt.timedelta(days=100), None, "código que não existe (exemplo)"])
    # Vencimento manual mais curto que o padrão
    registros.append(["1002", "ASO", hoje - dt.timedelta(days=150), hoje + dt.timedelta(days=12),
                      "médico reduziu para 6 meses (exemplo)"])

    cas = [
        ["Capacete classe B", "Fabricante A", "11111", hoje + dt.timedelta(days=400), "Sim"],
        ["Cinturão paraquedista", "Fabricante B", "22222", hoje + dt.timedelta(days=18), "Sim"],
        ["Talabarte duplo c/ absorvedor", "Fabricante B", "22223", hoje - dt.timedelta(days=10), "Sim"],
        ["Luva de vaqueta", "Fabricante C", "33333", hoje + dt.timedelta(days=45), "Sim"],
        ["Botina de segurança", "Fabricante D", "44444", hoje + dt.timedelta(days=700), "Sim"],
        ["Protetor auricular plug", "Fabricante E", "55555", hoje - dt.timedelta(days=60), "Não"],
        ["Óculos incolor", "Fabricante F", "66666", hoje + dt.timedelta(days=250), "Sim"],
    ]
    return funcionarios, registros, cas


def montar(caminho, exemplo, hoje):
    wb = Workbook()
    ws = wb.active
    ws.title = E.ABA_LEIAME
    for linha in LEIAME:
        ws.append([linha])
    ws["A1"].font = Font(bold=True, size=14)
    for celula in ("A3", "A11", "A14"):
        ws[celula].font = Font(bold=True)
    ws.column_dimensions["A"].width = 130

    ws = wb.create_sheet(E.ABA_ITENS)
    cabecalho(ws, E.COLS_ITENS, {1: 10, 2: 42, 3: 12, 4: 26, 5: 90})
    for item in E.ITENS_PADRAO:
        ws.append(list(item))

    ws = wb.create_sheet(E.ABA_FUNCOES)
    codigos = [cod for cod, *_ in E.ITENS_PADRAO]
    cabecalho(ws, ["Função"] + codigos, {1: 24})
    for funcao, exige in E.FUNCOES_PADRAO.items():
        ws.append([funcao] + [exige.get(cod) for cod in codigos])
    for col in range(2, len(codigos) + 2):
        ws.column_dimensions[get_column_letter(col)].width = 8
        for linha in range(2, 200):
            ws.cell(row=linha, column=col).alignment = Alignment(horizontal="center")

    funcionarios, registros, cas = dados_exemplo(hoje, random.Random(7)) if exemplo else ([], [], [])

    ws = wb.create_sheet(E.ABA_FUNCIONARIOS)
    cabecalho(ws, E.COLS_FUNCIONARIOS, {2: 28, 3: 22, 4: 30, 5: 26})
    for f in funcionarios:
        ws.append(f)
    formatar_datas(ws, [6])
    lista(ws, "G2:G2000", '"' + ",".join(E.SITUACOES) + '"')
    lista(ws, "C2:C2000", f"='{E.ABA_FUNCOES}'!$A$2:$A$200")

    ws = wb.create_sheet(E.ABA_REGISTROS)
    cabecalho(ws, E.COLS_REGISTROS, {3: 18, 4: 14, 5: 44})
    for r in registros:
        ws.append(r)
    formatar_datas(ws, [3, 4], ate_linha=5000)
    lista(ws, "B2:B5000", f"='{E.ABA_ITENS}'!$A$2:$A$100")

    ws = wb.create_sheet(E.ABA_CA)
    cabecalho(ws, E.COLS_CA, {1: 32, 2: 20, 4: 16})
    for c in cas:
        ws.append(c)
    formatar_datas(ws, [4], ate_linha=500)
    lista(ws, "E2:E500", '"Sim,Não"')

    wb.save(caminho)


def main():
    ap = argparse.ArgumentParser(description="Cria o modelo da planilha de controle de vencimentos de SST.")
    ap.add_argument("--saida", default=str(PASTA / "controle-sst.xlsx"))
    ap.add_argument("--exemplo", action="store_true", help="preenche com dados fictícios")
    ap.add_argument("--forcar", action="store_true", help="sobrescreve se já existir")
    args = ap.parse_args()

    caminho = Path(args.saida)
    if caminho.exists() and not args.forcar:
        sys.exit(f"Já existe {caminho}. Use --forcar para sobrescrever (cuidado com dados reais).")
    montar(caminho, args.exemplo, dt.date.today())
    print(f"Planilha criada: {caminho}" + (" (com dados fictícios)" if args.exemplo else ""))


if __name__ == "__main__":
    main()
